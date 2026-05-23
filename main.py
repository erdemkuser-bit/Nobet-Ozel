# -*- coding: utf-8 -*-
"""
main.py - Nobet Dagitim Otomasyonu GUI

Tkinter/ttk arayuzu, scheduler.py motoru ve openpyxl Excel cikti ureticisini
tek dosyada birlestirir.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from scheduler import (
    AY_ADLARI, NobetAtama, Personel,
    load_istisnai_tatiller,
    schedule_month, schedule_month_from_data,
)
from config import config
from auth import authenticate, is_password_set, setup_password
from data_manager import db
from audit_logger import audit

PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_EXCEL_PATH = config.excel_path

APP_NAME = "MUH-NET"
FULL_APP_NAME = "Muhabere Nöbet Etkileşim ve Takip Sistemi"

PRIMARY = "#0f3d67"
PRIMARY_DARK = "#08233d"
ACCENT = "#1d6fa5"
BG = "#edf2f7"
PANEL = "#ffffff"
TEXT = "#1f2933"
MUTED = "#66788a"
DANGER = "#b42318"
SUCCESS = "#0f7b45"
WARNING = "#b35c00"

KAT_ETIKET = {
    "Ozel3": "Ozel-3",
    "Ozel2": "Ozel-2",
    "Ozel1": "Ozel-1",
    "Cmt": "Cumartesi",
    "Pzr": "Pazar",
    "Cuma": "Cuma",
    "HI": "Hafta Ici",
}

TREE_TAG_COLORS = {
    "Ozel3": "#f8d7da",
    "Ozel2": "#ffe5cc",
    "Ozel1": "#fff3cd",
    "Cmt": "#d9ecff",
    "Pzr": "#e8e3f3",
    "Cuma": "#dff3e3",
    "HI": "#ffffff",
    "gap1": "#ffd6d6",
}

EXCEL_FILLS = {
    "Ozel3": "F4CCCC",
    "Ozel2": "FCE5CD",
    "Ozel1": "FFF2CC",
    "Cmt": "D9EAF7",
    "Pzr": "EADCF8",
    "Cuma": "D9EAD3",
    "HI": "FFFFFF",
}


def desktop_path() -> Path:
    """Windows/Turkish desktop variants; fallback to project dir."""
    candidates = [
        Path.home() / "Desktop",
        Path.home() / "Masaüstü",
        Path.home() / "Masaustu",
        PROJECT_DIR,
    ]
    for candidate in candidates:
        if candidate.exists() and candidate.is_dir():
            return candidate
    return PROJECT_DIR


def safe_person_name(p: Personel | None) -> str:
    if p is None:
        return ""
    return p.ad_soyad.strip() or p.sicil


def yedek_at(atama: NobetAtama, index: int) -> str:
    if len(atama.yedekler) <= index:
        return ""
    return safe_person_name(atama.yedekler[index])


class NobetApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(f"{APP_NAME} | {FULL_APP_NAME}")
        self.geometry("1180x760")
        self.minsize(1060, 680)
        self.configure(bg=BG)

        self.atamalar: list[NobetAtama] = []
        self.aktif_personel: list[Personel] = []
        self.atanmayan_gunler: list[tuple[date, str]] = []

        self.year_var = tk.StringVar(value=str(date.today().year))
        self.month_var = tk.StringVar(value=AY_ADLARI.get(date.today().month, "Mayis"))
        self.unit_var = tk.StringVar(value="BIRLIK ADI")
        self.excel_path_var = tk.StringVar(value=str(DEFAULT_EXCEL_PATH))
        self.status_var = tk.StringVar(value="Hazir")

        self._configure_styles()
        self._build_layout()
        self._write_log("Sistem hazir. Nöbet listesi oluşturmak için yıl, ay ve birlik adını girin.\n")
        self.bind("<F12>", self.gosterisli_imza_goster)
        # İlk açılış veya şifre kurulum/bütünlük kontrolleri
        self.after(400, self._startup_checks)

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        style.theme_use("clam")

        style.configure("TFrame", background=BG)
        style.configure("Panel.TFrame", background=PANEL)
        style.configure("Header.TFrame", background=PRIMARY)

        style.configure(
            "Header.TLabel",
            background=PRIMARY,
            foreground="white",
            font=("Segoe UI", 18, "bold"),
        )
        style.configure(
            "SubHeader.TLabel",
            background=PRIMARY,
            foreground="#dbeafe",
            font=("Segoe UI", 10),
        )
        style.configure("TLabel", background=BG, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("Panel.TLabel", background=PANEL, foreground=TEXT, font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=PANEL, foreground=MUTED, font=("Segoe UI", 9))

        style.configure(
            "Primary.TButton",
            background=PRIMARY,
            foreground="white",
            borderwidth=0,
            focusthickness=0,
            padding=(14, 8),
            font=("Segoe UI", 10, "bold"),
        )
        style.map("Primary.TButton", background=[("active", ACCENT), ("disabled", "#8aa1b4")])

        style.configure(
            "Secondary.TButton",
            background="#e2e8f0",
            foreground=TEXT,
            borderwidth=0,
            padding=(12, 8),
            font=("Segoe UI", 10),
        )
        style.map("Secondary.TButton", background=[("active", "#cbd5e1")])

        style.configure(
            "Treeview",
            background="white",
            foreground=TEXT,
            fieldbackground="white",
            rowheight=30,
            borderwidth=0,
            font=("Segoe UI", 10),
        )
        style.configure(
            "Treeview.Heading",
            background=PRIMARY_DARK,
            foreground="white",
            relief="flat",
            font=("Segoe UI", 10, "bold"),
            padding=(8, 8),
        )
        style.map("Treeview.Heading", background=[("active", PRIMARY)])

    def _build_layout(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(2, weight=1)

        header = ttk.Frame(self, style="Header.TFrame", padding=(22, 16))
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(0, weight=1)

        ttk.Label(header, text="NOBET DAGITIM OTOMASYONU", style="Header.TLabel").grid(
            row=0, column=0, sticky="w"
        )
        ttk.Label(
            header,
            text="JGY 64-4 (B) uyumlu vardiyali nobet planlama ve resmi Excel cizelge uretimi",
            style="SubHeader.TLabel",
        ).grid(row=1, column=0, sticky="w", pady=(3, 0))

        controls = ttk.Frame(self, style="Panel.TFrame", padding=(18, 14))
        controls.grid(row=1, column=0, sticky="ew", padx=16, pady=(14, 10))
        controls.columnconfigure(10, weight=1)

        ttk.Label(controls, text="Yil", style="Panel.TLabel").grid(row=0, column=0, sticky="w")
        year_spin = ttk.Spinbox(controls, from_=2020, to=2050, width=8, textvariable=self.year_var)
        year_spin.grid(row=1, column=0, sticky="w", padx=(0, 12), pady=(4, 0))

        ttk.Label(controls, text="Ay", style="Panel.TLabel").grid(row=0, column=1, sticky="w")
        month_combo = ttk.Combobox(
            controls,
            width=14,
            state="readonly",
            values=[AY_ADLARI[i] for i in range(1, 13)],
            textvariable=self.month_var,
        )
        month_combo.grid(row=1, column=1, sticky="w", padx=(0, 12), pady=(4, 0))

        ttk.Label(controls, text="Birlik Adi", style="Panel.TLabel").grid(row=0, column=2, sticky="w")
        ttk.Entry(controls, width=24, textvariable=self.unit_var).grid(
            row=1, column=2, sticky="w", padx=(0, 12), pady=(4, 0)
        )

        ttk.Label(controls, text="Excel Kaynagi", style="Panel.TLabel").grid(
            row=0, column=3, columnspan=3, sticky="w"
        )
        ttk.Entry(controls, width=42, textvariable=self.excel_path_var).grid(
            row=1, column=3, columnspan=3, sticky="ew", padx=(0, 8), pady=(4, 0)
        )
        ttk.Button(
            controls,
            text="Sec",
            style="Secondary.TButton",
            command=self._browse_excel,
        ).grid(row=1, column=6, padx=(0, 12), pady=(4, 0))

        self.generate_btn = ttk.Button(
            controls,
            text="Nobet Listesini Olustur",
            style="Primary.TButton",
            command=self.generate_schedule,
        )
        self.generate_btn.grid(row=1, column=7, padx=(0, 8), pady=(4, 0))

        self.export_btn = ttk.Button(
            controls,
            text="Excel Cizelgesi Olarak Aktar",
            style="Secondary.TButton",
            command=self.export_excel,
            state="disabled",
        )
        self.export_btn.grid(row=1, column=8, padx=(0, 8), pady=(4, 0))

        self.admin_btn = ttk.Button(
            controls,
            text="⚙ İdari Yönetim",
            style="Secondary.TButton",
            command=self._open_admin_panel,
        )
        self.admin_btn.grid(row=1, column=9, padx=(0, 8), pady=(4, 0))

        ttk.Label(controls, textvariable=self.status_var, style="Muted.TLabel").grid(
            row=1, column=10, sticky="e", padx=(14, 0), pady=(4, 0)
        )

        main_panel = ttk.Frame(self, style="Panel.TFrame", padding=(14, 14))
        main_panel.grid(row=2, column=0, sticky="nsew", padx=16, pady=(0, 10))
        main_panel.columnconfigure(0, weight=1)
        main_panel.rowconfigure(0, weight=1)

        columns = ("tarih", "kategori", "puan", "atanan", "yedek1", "yedek2", "gap")
        self.tree = ttk.Treeview(main_panel, columns=columns, show="headings", selectmode="browse")
        headings = {
            "tarih": "Tarih",
            "kategori": "Gun Kategorisi",
            "puan": "Puan",
            "atanan": "Atanan Personel",
            "yedek1": "Yedek 1",
            "yedek2": "Yedek 2",
            "gap": "Uygulanan Gap",
        }
        widths = {
            "tarih": 120,
            "kategori": 150,
            "puan": 70,
            "atanan": 250,
            "yedek1": 190,
            "yedek2": 190,
            "gap": 120,
        }
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(col, width=widths[col], anchor="center" if col in {"puan", "gap"} else "w")

        for tag, color in TREE_TAG_COLORS.items():
            self.tree.tag_configure(tag, background=color)

        tree_scroll_y = ttk.Scrollbar(main_panel, orient="vertical", command=self.tree.yview)
        tree_scroll_x = ttk.Scrollbar(main_panel, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=tree_scroll_y.set, xscrollcommand=tree_scroll_x.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        tree_scroll_y.grid(row=0, column=1, sticky="ns")
        tree_scroll_x.grid(row=1, column=0, sticky="ew")

        log_panel = ttk.Frame(self, style="Panel.TFrame", padding=(14, 10))
        log_panel.grid(row=3, column=0, sticky="ew", padx=16, pady=(0, 0))
        log_panel.columnconfigure(0, weight=1)

        ttk.Label(log_panel, text="Operasyonel Log ve Uyarilar", style="Panel.TLabel").grid(
            row=0, column=0, sticky="w", pady=(0, 6)
        )

        footer = ttk.Frame(self, style="TFrame")
        footer.grid(row=4, column=0, sticky="ew", padx=16, pady=(2, 6))
        footer.columnconfigure(0, weight=1)
        ttk.Label(
            footer,
            text=f"{APP_NAME}  —  {FULL_APP_NAME}    |    Yazılım Geliştirme: J.Uzm.Çvş.Mu. Erdem KÜSER  © 2026",
            font=("Segoe UI", 8, "italic"),
            foreground="#94a3b8",
            background=BG,
            anchor="e",
        ).grid(row=0, column=0, sticky="e")
        self.log_text = tk.Text(
            log_panel,
            height=8,
            wrap="word",
            bg="#f8fafc",
            fg=TEXT,
            insertbackground=TEXT,
            relief="flat",
            font=("Consolas", 10),
            padx=10,
            pady=8,
        )
        self.log_text.grid(row=1, column=0, sticky="ew")
        self.log_text.tag_configure("info", foreground=TEXT)
        self.log_text.tag_configure("success", foreground=SUCCESS, font=("Consolas", 10, "bold"))
        self.log_text.tag_configure("warning", foreground=WARNING, font=("Consolas", 10, "bold"))
        self.log_text.tag_configure("error", foreground=DANGER, font=("Consolas", 10, "bold"))

    def _browse_excel(self) -> None:
        path = filedialog.askopenfilename(
            title="Personel Excel dosyasini sec",
            initialdir=str(PROJECT_DIR),
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        )
        if path:
            self.excel_path_var.set(path)

    def _month_number(self) -> int:
        month_name = self.month_var.get()
        for number, name in AY_ADLARI.items():
            if name == month_name:
                return number
        return date.today().month

    def _write_log(self, message: str, tag: str = "info") -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", message, tag)
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _clear_log(self) -> None:
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

    def _set_busy(self, busy: bool) -> None:
        state = "disabled" if busy else "normal"
        self.generate_btn.configure(state=state)
        if not busy and self.atamalar:
            self.export_btn.configure(state="normal")
        elif busy:
            self.export_btn.configure(state="disabled")
        self.update_idletasks()

    def _populate_tree(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)

        for atama in self.atamalar:
            tag = "gap1" if atama.gap_kullanilan == 1 else atama.kategori
            self.tree.insert(
                "",
                "end",
                values=(
                    atama.tarih.strftime("%d.%m.%Y"),
                    KAT_ETIKET.get(atama.kategori, atama.kategori),
                    atama.puan,
                    safe_person_name(atama.atanan),
                    yedek_at(atama, 0),
                    yedek_at(atama, 1),
                    atama.gap_kullanilan,
                ),
                tags=(tag,),
            )

    def generate_schedule(self) -> None:
        # ── Yönetici yetkilendirmesi ──────────────────────────────────────────
        if not authenticate(self, "Nöbet listesi oluşturmak için yönetici şifresi gerekli:"):
            self._write_log("Yetkilendirme başarısız. İşlem iptal edildi.\n", "warning")
            return

        self._clear_log()
        self.status_var.set("Hesaplaniyor...")
        self.atamalar = []
        self.aktif_personel = []
        self.atanmayan_gunler = []
        self._populate_tree()
        self._set_busy(True)

        try:
            yil = int(self.year_var.get())
            ay = self._month_number()

            self._write_log(f"Nobet listesi olusturuluyor: {AY_ADLARI[ay]} {yil}\n")

            # DB'de personel varsa SQLite akışı; yoksa Excel fallback
            personel_list = db.get_scheduler_personel()
            if personel_list:
                self._write_log(
                    f"Veri kaynagi: secure_data.db ({len(personel_list)} personel)\n"
                )
                mazeretler    = db.get_scheduler_mazeretler(yil, ay)
                # Özel tatiller hâlâ Excel'den okunur (statik referans)
                excel_path = config.excel_path
                ozel_tatiller = (
                    load_istisnai_tatiller(str(excel_path))
                    if excel_path.exists() else {}
                )
                result = schedule_month_from_data(
                    yil, ay, personel_list, mazeretler, ozel_tatiller
                )
            else:
                # Fallback: Excel doğrudan (ilk kurulum tamamlanmamış olabilir)
                excel_path = Path(self.excel_path_var.get()).expanduser()
                if not excel_path.exists():
                    self._write_log(
                        f"HATA: Veritabanında personel yok ve Excel bulunamadı: {excel_path}\n",
                        "error",
                    )
                    self.status_var.set("Veri kaynagi yok")
                    return
                self._write_log(
                    f"Fallback: Excel kaynagi kullaniliyor: {excel_path}\n", "warning"
                )
                result = schedule_month(yil, ay, str(excel_path))

            if not isinstance(result, tuple) or len(result) != 3:
                self.atamalar = []
                self.aktif_personel = []
                self.atanmayan_gunler = []
                self._write_log("HATA: scheduler beklenen sonucu dondurmedi.\n", "error")
                self.status_var.set("Hesaplama hatasi")
                return

            self.atamalar, self.aktif_personel, self.atanmayan_gunler = result
            self._populate_tree()
            self._write_summary_log(yil, ay)
            self.export_btn.configure(state="normal" if self.atamalar else "disabled")
            self.status_var.set(f"{len(self.atamalar)} nobet listelendi")
            # Denetim kaydı
            audit.log(
                "Nöbet Listesi Oluşturuldu",
                details=(
                    f"{AY_ADLARI[ay]} {yil} | "
                    f"Toplam nöbet: {len(self.atamalar)} | "
                    f"Aktif personel: {len(self.aktif_personel)} | "
                    f"Atanmayan gün: {len(self.atanmayan_gunler)}"
                ),
            )
        except Exception as exc:
            self.atamalar = []
            self.export_btn.configure(state="disabled")
            self._write_log(f"HATA: {exc}\n", "error")
            self.status_var.set("Hata olustu")
            messagebox.showerror("Hata", str(exc))
        finally:
            self._set_busy(False)

    def _write_summary_log(self, yil: int, ay: int) -> None:
        gap1 = [a for a in self.atamalar if a.gap_kullanilan == 1]
        nobetci_sayisi = len({a.atanan.sicil for a in self.atamalar})

        self._write_log("\nOZET\n", "success")
        self._write_log(f"- Donem: {AY_ADLARI[ay]} {yil}\n")
        self._write_log(f"- Toplam nobet: {len(self.atamalar)}\n")
        self._write_log(f"- Nobet tutan personel: {nobetci_sayisi} / {len(self.aktif_personel)}\n")

        if gap1:
            dates = ", ".join(a.tarih.strftime("%d.%m") for a in gap1)
            self._write_log(
                f"- KRITIK UYARI: gap=1 kullanildi ({len(gap1)} gun): {dates}\n",
                "error",
            )
        else:
            self._write_log("- Gap=1 kritik sikisma yok.\n", "success")

        if self.atanmayan_gunler:
            self._write_log(f"- ATANAMAYAN GUN: {len(self.atanmayan_gunler)}\n", "error")
            for tarih, kategori in self.atanmayan_gunler:
                self._write_log(
                    f"  {tarih.strftime('%d.%m.%Y')} [{KAT_ETIKET.get(kategori, kategori)}]\n",
                    "error",
                )
        else:
            self._write_log("- Tum gunlere nobet atandi.\n", "success")

    def export_excel(self) -> None:
        if not self.atamalar:
            messagebox.showwarning("Uyari", "Once nobet listesi olusturulmalidir.")
            return

        try:
            yil = int(self.year_var.get())
            ay = self._month_number()
            default_filename = f"Nobet_Cizelgesi_{yil}_{ay:02d}.xlsx"
            selected_path = filedialog.asksaveasfilename(
                title="Nobet cizelgesini farkli kaydet",
                initialdir=str(desktop_path()),
                initialfile=default_filename,
                defaultextension=".xlsx",
                filetypes=[("Excel Files", "*.xlsx")],
            )

            if not selected_path:
                self._write_log("Aktarım kullanıcı tarafından iptal edildi\n", "warning")
                self.status_var.set("Aktarim iptal edildi")
                return

            output_path = self._build_export_workbook(yil, ay, Path(selected_path))
            self._write_log(f"Excel cizelgesi kaydedildi: {output_path}\n", "success")

            # DB güncelleme (birincil)
            db.update_personel_counters(self.aktif_personel)
            self._write_log("Veritabani (secure_data.db) puan sayaclari guncellendi.\n",
                            "success")

            # Excel güncelleme (geriye dönük uyumluluk / referans dosya)
            source_path = config.excel_path
            if source_path.exists():
                self._update_source_personel(source_path)
                self._write_log(
                    f"Referans Excel de guncellendi: {source_path}\n", "success"
                )

            audit.log(
                "Excel Çizelgesi Aktarıldı",
                details=f"Dosya: {output_path} | "
                        f"{AY_ADLARI[ay]} {yil} | {len(self.atamalar)} nöbet",
            )
            self.status_var.set("Excel aktarildi")
            messagebox.showinfo(
                "Başarılı",
                f"Resmi çizelge kaydedildi:\n{output_path}\n\n"
                "Veritabanı puan sayaçları güncellendi.",
            )
        except Exception as exc:
            self._write_log(f"EXCEL HATASI: {exc}\n", "error")
            self.status_var.set("Excel hatasi")
            messagebox.showerror("Excel Hatasi", str(exc))

    def _build_export_workbook(self, yil: int, ay: int, output_path: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Nobet_Cizelgesi"

        thin = Side(style="thin", color="B8C2CC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="0F3D67")
        sub_fill = PatternFill("solid", fgColor="D9EAF7")
        white_font = Font(color="FFFFFF", bold=True)
        title_font = Font(color="0F3D67", bold=True, size=16)
        bold_font = Font(bold=True, color="1F2933")
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        ws.merge_cells("A1:G1")
        ws["A1"] = "NOBET CIZELGESI"
        ws["A1"].font = title_font
        ws["A1"].alignment = center

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Birlik Adi: {self.unit_var.get().strip() or 'BIRLIK ADI'}"
        ws["A2"].alignment = center
        ws["A2"].font = bold_font

        ws.merge_cells("A3:G3")
        ws["A3"] = f"Donem: {AY_ADLARI[ay]} {yil}    |    Olusturma Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
        ws["A3"].alignment = center
        ws["A3"].font = Font(color="4A5568", italic=True)

        headers = [
            "Tarih",
            "Gun Kategorisi",
            "Puan",
            "Atanan Personel",
            "Yedek 1",
            "Yedek 2",
            "Uygulanan Gap",
        ]
        header_row = 6
        for col_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_index, value=header)
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = center
            cell.border = border

        for row_index, atama in enumerate(self.atamalar, start=header_row + 1):
            row_values = [
                atama.tarih.strftime("%d.%m.%Y"),
                KAT_ETIKET.get(atama.kategori, atama.kategori),
                atama.puan,
                safe_person_name(atama.atanan),
                yedek_at(atama, 0),
                yedek_at(atama, 1),
                atama.gap_kullanilan,
            ]
            fill_color = "FFD6D6" if atama.gap_kullanilan == 1 else EXCEL_FILLS.get(atama.kategori, "FFFFFF")
            fill = PatternFill("solid", fgColor=fill_color)

            for col_index, value in enumerate(row_values, start=1):
                cell = ws.cell(row=row_index, column=col_index, value=value)
                cell.fill = fill
                cell.border = border
                cell.alignment = center if col_index in {1, 2, 3, 7} else left
                if atama.gap_kullanilan == 1:
                    cell.font = Font(color="9B1C1C", bold=True)

        summary_start = header_row + len(self.atamalar) + 3
        ws.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=7)
        ws.cell(row=summary_start, column=1, value="Onay Makami ve Imza Bolumu")
        ws.cell(row=summary_start, column=1).fill = sub_fill
        ws.cell(row=summary_start, column=1).font = bold_font
        ws.cell(row=summary_start, column=1).alignment = center

        approval_row = summary_start + 2
        approvals = [("Hazirlayan", 1, 2), ("Kontrol Eden", 3, 5), ("Onay", 6, 7)]
        for label, start_col, end_col in approvals:
            ws.merge_cells(start_row=approval_row, start_column=start_col, end_row=approval_row, end_column=end_col)
            ws.cell(row=approval_row, column=start_col, value=label)
            ws.cell(row=approval_row, column=start_col).alignment = center
            ws.cell(row=approval_row, column=start_col).font = bold_font
            ws.cell(row=approval_row, column=start_col).border = border

            ws.merge_cells(start_row=approval_row + 1, start_column=start_col, end_row=approval_row + 4, end_column=end_col)
            sign_cell = ws.cell(row=approval_row + 1, column=start_col, value="\n\nAd Soyad / Rutbe / Imza")
            sign_cell.alignment = center
            sign_cell.border = border

        widths = [14, 18, 10, 28, 24, 24, 16]
        for col_index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        for row in range(1, approval_row + 5):
            ws.row_dimensions[row].height = 24
        ws.row_dimensions[1].height = 34
        ws.row_dimensions[approval_row + 1].height = 70

        ws.freeze_panes = "A7"
        ws.auto_filter.ref = f"A{header_row}:G{header_row + len(self.atamalar)}"

        self._add_personel_puan_sheet(wb)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    def _add_personel_puan_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Guncel_Puanlar")

        thin = Side(style="thin", color="B8C2CC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="0F3D67")
        white_font = Font(color="FFFFFF", bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        headers = [
            "Sicil/Id",
            "Rutbe",
            "Ad Soyad",
            "Toplam Puan",
            "Hafta Ici",
            "Cuma",
            "Cumartesi",
            "Pazar",
            "Ozel Nobet",
            "Yedek",
            "Ay Ici Nobet",
            "Son Nobet Tarihi",
        ]

        for col_index, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_index, value=header)
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = center
            cell.border = border

        for row_index, personel in enumerate(self.aktif_personel, start=2):
            values = [
                personel.sicil,
                personel.rutbe,
                personel.ad_soyad,
                personel.toplam_puan,
                personel.hafta_ici,
                personel.cuma,
                personel.cumartesi,
                personel.pazar,
                personel.ozel_nobet,
                personel.yedek,
                personel.ay_ici_nobet,
                personel.son_nobet_tarihi.strftime("%d.%m.%Y") if personel.son_nobet_tarihi else "",
            ]
            for col_index, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=col_index, value=value)
                cell.border = border
                cell.alignment = left if col_index in {1, 2, 3} else center

        widths = [14, 18, 26, 14, 12, 10, 12, 10, 12, 10, 12, 16]
        for col_index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:L{max(len(self.aktif_personel) + 1, 1)}"

    def _update_source_personel(self, source_path: Path) -> None:
        """
        Orijinal personel_listesi.xlsx dosyasinin Personel_Havuzu sayfasini
        guncel puan ve sayaclarla in-place gunceller.

        Personel_Havuzu sutun duzeni (create_excel_template.py ile eslesen):
          A=Sicil  B=Rutbe  C=AdSoyad  D=KidemYili  E=MuafiyetTuru  F=GrupAdi
          G=ToplamPuan  H=HaftaIci  I=Cuma  J=Cumartesi  K=Pazar  L=OzelNobet
          M=SonNobetTarihi

        Satirlar: 1=ana baslik, 2=sutun basliklari, 3+ = veri.
        Sadece G–M sutunlari (indeks 7–13) guncellenir; diger alanlar dokunulmaz.
        """
        if not source_path.exists():
            raise FileNotFoundError(
                f"Kaynak dosya bulunamadi: {source_path}\n"
                "Kaynak dosya guncellenemedi; cizelge dosyasi yine de kaydedildi."
            )

        # Guncel personel verilerini sicil'e gore bir haritaya al
        puan_haritasi: dict[str, Personel] = {
            p.sicil: p for p in self.aktif_personel
        }

        from openpyxl import load_workbook as _load_wb
        wb = _load_wb(source_path)
        ws = wb["Personel_Havuzu"]

        for row in ws.iter_rows(min_row=3):
            sicil_cell = row[0]
            if not sicil_cell.value:
                continue
            sicil = str(sicil_cell.value).strip()
            if len(sicil) > 20 or sicil not in puan_haritasi:
                continue

            p = puan_haritasi[sicil]

            # G=ToplamPuan (col index 6)
            row[6].value = p.toplam_puan
            # H=HaftaIci (col index 7)
            row[7].value = p.hafta_ici
            # I=Cuma (col index 8)
            row[8].value = p.cuma
            # J=Cumartesi (col index 9)
            row[9].value = p.cumartesi
            # K=Pazar (col index 10)
            row[10].value = p.pazar
            # L=OzelNobet (col index 11)
            row[11].value = p.ozel_nobet
            # M=SonNobetTarihi (col index 12) — her zaman DD.MM.YYYY metin olarak yaz.
            # date/datetime nesnesi yazılırsa openpyxl Excel seri numarasına çevirir;
            # sonraki okumada tip uyuşmazlığı yaşanmaması için saf metin tercih edilir.
            row[12].value = (
                p.son_nobet_tarihi.strftime("%d.%m.%Y")
                if p.son_nobet_tarihi is not None
                else ""
            )
            row[12].number_format = "@"   # sütunu "Metin" olarak işaretle

        wb.save(source_path)
        wb.close()

    # ── Startup ve Admin Panel Yardımcıları ───────────────────────────────────

    def _startup_checks(self) -> None:
        """Uygulama açıldıktan kısa süre sonra çalışır."""
        # İlk kurulum: şifre henüz yok
        if not is_password_set():
            messagebox.showinfo(
                "İlk Kurulum",
                f"Hoş geldiniz! {APP_NAME} ilk kez başlatılıyor.\n\n"
                "Lütfen bir yönetici şifresi oluşturun.",
                parent=self,
            )
            if setup_password(self):
                self._do_first_run_import()
            else:
                self._write_log(
                    "UYARI: Şifre oluşturulmadı. Kritik işlemler kısıtlı olacak.\n",
                    "warning",
                )
            return

        # Bütünlük kontrolü (DB kuruluysa)
        self._run_startup_integrity_check()

    def _do_first_run_import(self) -> None:
        """İlk kurulumda Excel'den personeli DB'ye aktarır."""
        excel = config.excel_path
        if not excel.exists():
            self._write_log(
                f"Excel referans dosyası bulunamadı: {excel}\n"
                "İdari Yönetim Paneli üzerinden personel ekleyebilirsiniz.\n",
                "warning",
            )
            return
        added, skipped, errors = db.sync_from_excel(str(excel))
        self._write_log(
            f"İlk kurulum — Excel'den içe aktarım: "
            f"{added} eklendi, {skipped} atlandı, {errors} hata.\n",
            "success" if errors == 0 else "warning",
        )
        config.mark_configured()

    def _run_startup_integrity_check(self) -> None:
        """DB ile Excel arasındaki personel farkını kontrol eder."""
        excel = config.excel_path
        if not excel.exists():
            return
        warns = db.integrity_check(str(excel))
        if warns:
            self._write_log(
                "BÜTÜNLÜK UYARISI: Excel ile DB arasında fark tespit edildi!\n",
                "warning",
            )
            for w in warns:
                self._write_log(f"  • {w}\n", "warning")
            # Kritik denetim kaydı
            audit.log_critical(
                "Açılış Bütünlük Kontrolü — Uyumsuzluk",
                details=f"{len(warns)} uyarı: " + " | ".join(warns[:5])
                        + (" ..." if len(warns) > 5 else ""),
            )
            messagebox.showwarning(
                "Veri Bütünlüğü Uyarısı",
                "Excel dosyası ile veritabanı arasında personel uyumsuzluğu:\n\n"
                + "\n".join(f"• {w}" for w in warns)
                + "\n\nAyrıntılar için İdari Yönetim Paneli → Sistem Denetim sekmesine bakın.",
                parent=self,
            )
        else:
            audit.log("Açılış Bütünlük Kontrolü",
                      details="DB ile Excel uyumlu — sorun yok.")

    def _open_admin_panel(self) -> None:
        """Yönetici yetkilendirmesi ile İdari Yönetim Paneli'ni açar."""
        if not authenticate(self, "İdari Yönetim Paneli için şifre gerekli:"):
            self._write_log("İdari panel erişimi reddedildi.\n", "warning")
            return
        from gui_manager import AdminPanel
        AdminPanel(self)

    def gosterisli_imza_goster(self, _event=None) -> None:
        """F12 — koyu temalı geliştirici bilgi paneli."""
        # ── Pencere ────────────────────────────────────────────────────────────
        panel = tk.Toplevel(self)
        panel.title("Sistem Gelistirici Bilgisi")
        panel.geometry("560x400")
        panel.resizable(False, False)
        panel.configure(bg="#0f172a")
        panel.grab_set()
        panel.focus_set()

        # Ekrana ortala
        panel.update_idletasks()
        x = self.winfo_x() + (self.winfo_width() - 560) // 2
        y = self.winfo_y() + (self.winfo_height() - 400) // 2
        panel.geometry(f"560x400+{x}+{y}")

        # ── Başlık bandı ───────────────────────────────────────────────────────
        title_bar = tk.Frame(panel, bg="#1e293b", height=70)
        title_bar.pack(fill="x")
        title_bar.pack_propagate(False)

        tk.Label(
            title_bar,
            text=f"🚀  {APP_NAME}  —  v1.0.0",
            bg="#1e293b",
            fg="#f59e0b",
            font=("Segoe UI", 15, "bold"),
            anchor="center",
        ).pack(pady=(10, 0))

        tk.Label(
            title_bar,
            text=FULL_APP_NAME,
            bg="#1e293b",
            fg="#94a3b8",
            font=("Segoe UI", 9, "italic"),
            anchor="center",
        ).pack()

        # ── İçerik alanı ──────────────────────────────────────────────────────
        content = tk.Frame(panel, bg="#0f172a")
        content.pack(fill="both", expand=True, padx=24, pady=(18, 12))

        # ── Sol: dairesel askeri amblem ────────────────────────────────────────
        canvas = tk.Canvas(
            content, width=140, height=160,
            bg="#0f172a", highlightthickness=0,
        )
        canvas.grid(row=0, column=0, rowspan=4, sticky="n", padx=(0, 24), pady=4)

        # Dış halka
        canvas.create_oval(10, 10, 130, 130, outline="#f59e0b", width=3)
        # İç halka
        canvas.create_oval(24, 24, 116, 116, outline="#475569", width=1)
        # Dolgu çemberi
        canvas.create_oval(30, 30, 110, 110, fill="#1e293b", outline="")

        # Yıldırım gövdesi (çizgi segmentleri)
        bolt = [70, 38, 54, 70, 66, 70, 50, 102, 70, 68, 58, 68, 74, 38]
        canvas.create_polygon(bolt, fill="#f59e0b", outline="#fbbf24", width=1)

        # Çevre nokta bezemeleri (altı nokta)
        import math
        for i in range(6):
            angle = math.radians(60 * i - 90)
            cx = 70 + 52 * math.cos(angle)
            cy = 70 + 52 * math.sin(angle)
            canvas.create_oval(cx - 4, cy - 4, cx + 4, cy + 4, fill="#f59e0b", outline="")

        # MUHABERE yazısı
        canvas.create_text(
            70, 148,
            text="MUHABERE",
            fill="#94a3b8",
            font=("Segoe UI", 9, "bold"),
            anchor="center",
        )

        # ── Sağ: metin bilgileri ───────────────────────────────────────────────
        right = tk.Frame(content, bg="#0f172a")
        right.grid(row=0, column=1, sticky="nw")

        tk.Label(
            right,
            text="J.Uzm.Çvş.Mu. Erdem KÜSER",
            bg="#0f172a",
            fg="#f59e0b",
            font=("Segoe UI", 15, "bold"),
            anchor="w",
        ).pack(anchor="w", pady=(0, 12))

        bilgiler = [
            ("Geliştirme Tarihi", "Mayıs 2026"),
            ("Yönetmelik Uyumu",  "JGY 64-4 (B)"),
            ("Sistem Altyapısı",  "Python 3.12 / Tkinter / openpyxl"),
        ]
        for etiket, deger in bilgiler:
            satir = tk.Frame(right, bg="#0f172a")
            satir.pack(anchor="w", pady=3)
            tk.Label(
                satir,
                text=f"{etiket}: ",
                bg="#0f172a",
                fg="#94a3b8",
                font=("Segoe UI", 10),
            ).pack(side="left")
            tk.Label(
                satir,
                text=deger,
                bg="#0f172a",
                fg="#e2e8f0",
                font=("Segoe UI", 10, "bold"),
            ).pack(side="left")

        # Ayırıcı çizgi
        sep = tk.Frame(right, bg="#334155", height=1)
        sep.pack(fill="x", pady=(14, 10))

        tk.Label(
            right,
            text="© 2026  Tüm hakları saklıdır.",
            bg="#0f172a",
            fg="#475569",
            font=("Segoe UI", 9, "italic"),
            anchor="w",
        ).pack(anchor="w")

        # ── Tamam butonu ───────────────────────────────────────────────────────
        btn_frame = tk.Frame(panel, bg="#0f172a")
        btn_frame.pack(fill="x", padx=24, pady=(0, 18))

        tamam = tk.Button(
            btn_frame,
            text="TAMAM",
            bg="#1e40af",
            fg="white",
            activebackground="#2563eb",
            activeforeground="white",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            cursor="hand2",
            padx=32,
            pady=8,
            command=panel.destroy,
        )
        tamam.pack(side="right")
        panel.bind("<Return>", lambda _e: panel.destroy())
        panel.bind("<Escape>", lambda _e: panel.destroy())


def main() -> None:
    app = NobetApp()
    app.mainloop()


if __name__ == "__main__":
    main()
